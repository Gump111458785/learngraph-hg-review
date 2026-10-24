import json

import pytest
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject, NameObject, NumberObject, StreamObject

from learngraph_review.cli import main
from learngraph_review.errors import ExtractionError
from learngraph_review.extract import extract_file, material_from_text


@pytest.mark.parametrize("encoding", ["utf-8", "utf-8-sig", "utf-16"])
def test_text_encoding_chunking_and_stable_ids(tmp_path, encoding):
    text = "第一行\n" + "可变序列。" * 500
    path = tmp_path / "资料.md"
    path.write_bytes(text.encode(encoding))
    material = extract_file(path)
    assert "".join(b.text for b in material.blocks) == text
    assert material.blocks[1].source_start == 1800
    assert material.id == extract_file(path).id
    assert material.source == "资料.md"


def test_docx_preserves_paragraph_table_order(tmp_path):
    doc = Document()
    doc.add_paragraph("正文第一段")
    cells = doc.add_table(rows=1, cols=2).rows[0].cells
    cells[0].text = "列表"
    cells[1].text = "可变"
    doc.add_paragraph("正文最后一段")
    path = tmp_path / "资料.docx"
    doc.save(path)
    material = extract_file(path)
    assert [b.text for b in material.blocks] == ["正文第一段", "列表\t可变", "正文最后一段"]
    assert [b.paragraph for b in material.blocks] == [1, 2, 3]


def write_pdf(path, *, with_text):
    writer = PdfWriter()
    for number in (1, 2):
        page = writer.add_blank_page(width=300, height=300)
        if with_text:
            font = DictionaryObject(
                {
                    NameObject("/Type"): NameObject("/Font"),
                    NameObject("/Subtype"): NameObject("/Type1"),
                    NameObject("/BaseFont"): NameObject("/Helvetica"),
                }
            )
            page[NameObject("/Resources")] = DictionaryObject(
                {
                    NameObject("/Font"): DictionaryObject({NameObject("/F1"): font}),
                }
            )
            stream = StreamObject()
            stream.set_data(f"BT /F1 12 Tf 20 250 Td (Page {number}) Tj ET".encode())
            page[NameObject("/Contents")] = writer._add_object(stream)
            page[NameObject("/Rotate")] = NumberObject(0)
    writer.write(path)


def test_pdf_text_and_page_positions(tmp_path):
    path = tmp_path / "two-pages.pdf"
    write_pdf(path, with_text=True)
    material = extract_file(path)
    assert [b.page for b in material.blocks] == [1, 2]
    assert material.blocks[0].text.strip() == "Page 1"


def test_pdf_without_text_rejects_instead_of_silently_approving(tmp_path):
    path = tmp_path / "scan.pdf"
    write_pdf(path, with_text=False)
    with pytest.raises(ExtractionError, match="OCR"):
        extract_file(path)


def test_pptx_text_and_slide_positions(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    for text in ("Python 列表可变", "元组不可变"):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        frame = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text_frame
        frame.text = text
    path = tmp_path / "slides.pptx"
    presentation.save(path)
    material = extract_file(path)
    assert [b.page for b in material.blocks] == [1, 2]
    assert [b.text for b in material.blocks] == ["Python 列表可变", "元组不可变"]


@pytest.mark.parametrize(
    "name,raw",
    [
        ("empty.txt", b""),
        ("binary.txt", b"\x00\x00"),
        ("bad.pdf", b"not a pdf"),
        ("bad.docx", b"not a zip"),
        ("script.py", b"print(1)"),
    ],
)
def test_unsupported_corrupt_or_empty_files_fail(tmp_path, name, raw):
    path = tmp_path / name
    path.write_bytes(raw)
    with pytest.raises(ExtractionError):
        extract_file(path)


def test_large_inputs_fail_without_truncation(tmp_path):
    path = tmp_path / "large.txt"
    path.write_bytes(b"a" * (20 * 1024 * 1024 + 1))
    with pytest.raises(ExtractionError, match="20 MiB"):
        extract_file(path)
    with pytest.raises(ExtractionError, match="160000"):
        material_from_text("字" * 160001, title="too large")


def test_cli_offline_extract_schema_and_no_overwrite(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    source = tmp_path / "note.txt"
    source.write_text("列表是可变序列。", encoding="utf-8")
    target = tmp_path / "material.json"
    assert main(["extract", str(source), "--id", "s1", "--out", str(target)]) == 0
    assert json.loads(target.read_text())["id"] == "s1"
    before = target.read_bytes()
    assert main(["extract", str(source), "--out", str(target)]) == 1
    assert target.read_bytes() == before
    schema = tmp_path / "schemas.json"
    assert main(["schemas", "--out", str(schema)]) == 0
    assert set(json.loads(schema.read_text())) == {
        "Context",
        "Material",
        "LessonDraft",
        "HReport",
        "GReport",
    }


def test_cli_handoff_and_g_revision_loop(
    tmp_path, monkeypatch, context, material, sources, bad_lesson, fixed_lesson
):
    from conftest import FixtureGateway, contradiction, g_result, h_result

    import learngraph_review.cli as cli
    from learngraph_review import ReviewService

    monkeypatch.chdir(tmp_path)
    service = ReviewService(
        FixtureGateway(
            h_result(material, ["accept", "reject"]),
            g_result(context, [contradiction(bad_lesson, sources[0])]),
            g_result(context),
        )
    )
    monkeypatch.setattr(cli, "ReviewService", lambda: service)
    for name, value in {
        "context": context.model_dump(),
        "material": material.model_dump(),
        "sources": [s.model_dump() for s in sources],
        "bad": bad_lesson.model_dump(),
        "fixed": fixed_lesson.model_dump(),
    }.items():
        (tmp_path / f"{name}.json").write_text(json.dumps(value), encoding="utf-8")
    assert (
        main(["h", "--context", "context.json", "--material", "material.json", "--out", "h.json"])
        == 2
    )
    assert (
        main(
            [
                "handoff",
                "--context",
                "context.json",
                "--material",
                "material.json",
                "--report",
                "h.json",
                "--out",
                "approved.json",
            ]
        )
        == 0
    )
    assert len(json.loads((tmp_path / "approved.json").read_text())[0]["blocks"]) == 1
    shared = ["--context", "context.json", "--sources", "sources.json"]
    assert main(["g", *shared, "--lesson", "bad.json", "--out", "g1.json"]) == 2
    original = (tmp_path / "g1.json").read_bytes()
    assert (
        main(["g", *shared, "--lesson", "fixed.json", "--previous", "g1.json", "--out", "g2.json"])
        == 0
    )
    assert json.loads((tmp_path / "g2.json").read_text())["status"] == "pass"
    assert original == (tmp_path / "g1.json").read_bytes()
