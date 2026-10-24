from __future__ import annotations

import io
import zipfile
from pathlib import Path

from .errors import ExtractionError
from .models import Material, TextBlock, digest

MAX_FILE_BYTES = 20 * 1024 * 1024
MAX_EXTRACTED_CHARS = 160_000
BLOCK_CHARS = 1800


def _blocks(
    text: str,
    blocks: list[TextBlock],
    *,
    page: int | None = None,
    paragraph: int | None = None,
) -> None:
    for start in range(0, len(text), BLOCK_CHARS):
        part = text[start : start + BLOCK_CHARS]
        if part.strip():
            blocks.append(
                TextBlock(
                    id=f"b{len(blocks) + 1:04d}",
                    text=part,
                    page=page,
                    paragraph=paragraph,
                    source_start=start,
                )
            )


def material_from_text(
    text: str,
    *,
    title: str,
    source: str = "provided-text",
    material_id: str | None = None,
) -> Material:
    if not text.strip():
        raise ExtractionError("文件没有可提取文本；扫描件请先进行 OCR")
    if len(text) > MAX_EXTRACTED_CHARS:
        raise ExtractionError("提取文本超过 160000 字符，请拆分文件后再审查")
    if "\x00" in text:
        raise ExtractionError("文件含 NUL 字符，可能不是文本文件或编码不受支持")
    blocks: list[TextBlock] = []
    _blocks(text, blocks)
    return Material(
        id=material_id or f"src-{digest([title, source, text])[:16]}",
        title=title,
        source=source,
        blocks=blocks,
    )


def extract_file(path: str | Path, *, material_id: str | None = None) -> Material:
    """Extract locally only. No file upload, URL fetch, OCR, or executable document content."""
    path = Path(path)
    if not path.is_file():
        raise ExtractionError("资料文件不存在或不是普通文件")
    if path.stat().st_size > MAX_FILE_BYTES:
        raise ExtractionError("文件超过 20 MiB，请拆分后再上传")
    raw = path.read_bytes()
    if len(raw) > MAX_FILE_BYTES:
        raise ExtractionError("文件超过 20 MiB")
    suffix = path.suffix.lower()
    notes: list[str] = []
    blocks: list[TextBlock] = []
    extracted_chars = 0
    try:
        if suffix in {".txt", ".md"}:
            encoding = "utf-16" if raw.startswith((b"\xff\xfe", b"\xfe\xff")) else "utf-8-sig"
            material = material_from_text(
                raw.decode(encoding),
                title=path.name,
                source=path.name,
                material_id=material_id,
            )
            return material
        if suffix == ".pdf":
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(raw))
            if reader.is_encrypted:
                raise ExtractionError("不支持加密 PDF，请先解密后再上传")
            if len(reader.pages) > 100:
                raise ExtractionError("PDF 超过 100 页，请拆分后再上传")
            for page_no, page in enumerate(reader.pages, 1):
                text = page.extract_text() or ""
                if not text.strip():
                    raise ExtractionError(f"PDF 第 {page_no} 页没有可提取文本，请 OCR 或移除空白页")
                extracted_chars += len(text)
                _check_size(extracted_chars)
                _blocks(text, blocks, page=page_no)
            notes = ["仅提取 PDF 文本层；未核验图片、公式图像或视觉排版；不执行 OCR。"]
        elif suffix == ".docx":
            from docx import Document
            from docx.table import Table

            with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                if sum(x.file_size for x in archive.infolist()) > 50 * 1024 * 1024:
                    raise ExtractionError("DOCX 解压后超过 50 MiB")
            document = Document(io.BytesIO(raw))
            ordinal = 0
            for item in document.iter_inner_content():
                texts = (
                    ["\t".join(c.text for c in row.cells) for row in item.rows]
                    if isinstance(item, Table)
                    else [item.text]
                )
                for text in texts:
                    ordinal += 1
                    extracted_chars += len(text)
                    _check_size(extracted_chars)
                    _blocks(text, blocks, paragraph=ordinal)
            notes = [
                "提取正文段落和表格行；paragraph 为正文/表格行顺序编号。",
                "不含页眉页脚、批注、文本框及图片内文字；未检查视觉排版。",
            ]
        elif suffix == ".pptx":
            from pptx import Presentation

            with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                if sum(x.file_size for x in archive.infolist()) > 50 * 1024 * 1024:
                    raise ExtractionError("PPTX 解压后超过 50 MiB")
            presentation = Presentation(io.BytesIO(raw))
            for slide_no, slide in enumerate(presentation.slides, 1):
                texts = []
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        texts.append(shape.text_frame.text)
                    if shape.has_table:
                        texts.extend(
                            "\t".join(c.text for c in row.cells) for row in shape.table.rows
                        )
                text = "\n".join(texts)
                if not text.strip():
                    raise ExtractionError(f"PPTX 第 {slide_no} 页无可提取文本，请先 OCR 或移除空页")
                extracted_chars += len(text)
                _check_size(extracted_chars)
                _blocks(text, blocks, page=slide_no)
            notes = ["page 为幻灯片编号；只提取文本框和表格，不含图片文字、讲者备注或分组内文本。"]
        else:
            raise ExtractionError("只支持 .txt、.md、可提取文本的 .pdf、.docx 和 .pptx")
    except ExtractionError:
        raise
    except Exception as exc:
        message = f"文件解析失败（{type(exc).__name__}），请检查文件格式与编码"
        raise ExtractionError(message) from exc
    if not blocks:
        raise ExtractionError("文件没有可提取文本；扫描件请先进行 OCR")
    return Material(
        id=material_id or f"src-{digest([path.name, raw.hex()])[:16]}",
        title=path.name,
        source=path.name,
        blocks=blocks,
        extraction_notes=notes,
    )


def _check_size(extracted_chars: int) -> None:
    if extracted_chars > MAX_EXTRACTED_CHARS:
        raise ExtractionError("提取文本超过 160000 字符，请拆分文件")
